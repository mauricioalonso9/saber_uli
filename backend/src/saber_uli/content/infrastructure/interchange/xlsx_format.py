"""Plantilla XLSX de ítems de selección (research R-04; FR-029, FR-037).

Una fila por ítem de LC, RC, CC o Inglés partes 2 a 7, con su contexto de texto. Los encabezados
de la hoja `items` son los nombres del formato v1, así cada fila se valida con los mismos modelos
que el JSON y los errores hablan de los mismos campos. Varias filas con el mismo código de
contexto lo comparten: manda la primera; un texto distinto en otra fila es un error de esa fila.

Defensas antes de abrir el archivo: tamaño (10 MB), contenido descomprimido (50 MB, contra
bombas zip), ZIP con `[Content_Types].xml` y sin macros (un `.xlsm` se rechaza). openpyxl usa
defusedxml (instalado) contra entidades externas. Las fórmulas se leen como el texto que son: no
se evalúan, y `data_only` no sirve porque entregaría el último valor que guardó Excel.
"""

import io
import zipfile
from collections.abc import Iterable, Iterator
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from saber_uli.content.application.bundle import (
    MAX_IMPORT_BYTES,
    MAX_IMPORT_ITEMS,
    BatchTooLargeError,
    Bundle,
    ContextSpec,
    EntryError,
    EntryKind,
    ItemSpec,
    UnreadableFileError,
    UnsupportedFileFormatError,
)
from saber_uli.content.domain.context import ContextKind
from saber_uli.content.domain.errors import FileTooLargeError
from saber_uli.content.domain.item_version import ItemKind
from saber_uli.content.domain.taxonomy import TaxonomyNode, TaxonomyNodeKind
from saber_uli.content.domain.validation import ContentIssue, ValidationRule
from saber_uli.content.infrastructure.interchange import mapping
from saber_uli.content.infrastructure.interchange import models as m
from saber_uli.content.infrastructure.interchange.json_format import (
    SCHEMA_VERSION,
    issues_from,
)

__all__ = [
    "COLUMNS",
    "MAX_FILE_BYTES",
    "MAX_UNCOMPRESSED_BYTES",
    "OPTIONAL_COLUMNS",
    "UnsupportedFileFormatError",
    "XlsxExportWriter",
    "read_xlsx",
    "write_template",
]

MAX_FILE_BYTES = MAX_IMPORT_BYTES
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
SHEET = "items"
KEYS = "ABCD"
TEMPLATE_ROWS = 1000

_TAGS = ("competencia", "afirmacion", "evidencia", "componente", "situacion")
COLUMNS: tuple[str, ...] = (
    "id",
    "modulo",
    "parte",
    *_TAGS,
    "dificultad",
    "origen",
    "contexto_id",
    "contexto_titulo",
    "contexto_tipo_texto",
    "contexto_texto",
    "enunciado",
    *(name for key in KEYS for name in (f"opcion_{key}", f"retroalimentacion_{key}")),
    "clave_correcta",
    "explicacion",
)
# Columnas que una plantilla anterior puede no tener.
OPTIONAL_COLUMNS: tuple[str, ...] = ("contexto_subtipo",)

_INSTRUCTIONS: tuple[tuple[str, str], ...] = (
    ("id", "Código del ítem, por ejemplo LC-0101. Si ya existe, se crea una versión nueva."),
    ("modulo", "LC, RC, CC o ING. Comunicación escrita va en el formato JSON."),
    ("parte", "Solo Inglés: 2 a 7. La parte 1 (emparejamiento) va en el formato JSON."),
    ("competencia … situacion", "Códigos de la hoja «listas», según el módulo."),
    ("dificultad", "baja, media o alta."),
    ("origen", "docente o ia. Los ítems de origen ia entran En revisión."),
    ("contexto_id", "Código del contexto. Vacío si el ítem no tiene contexto."),
    (
        "contexto_titulo / contexto_texto",
        "Solo la primera fila de cada contexto nuevo. Las "
        "demás filas del mismo contexto los dejan vacíos o los repiten iguales. Para usar un "
        "contexto que ya está en el banco, deje ambos vacíos.",
    ),
    ("contexto_tipo_texto", "Solo Lectura crítica: el tipo de texto del contexto."),
    ("contexto_subtipo", "Opcional: el subtipo del texto (por ejemplo, columna de opinión)."),
    ("enunciado", "La pregunta. Admite Markdown."),
    ("opcion_A … opcion_D", "Texto de cada opción (Inglés partes 2 a 5: A a C)."),
    ("retroalimentacion_A … D", "Por qué es correcta, o el error que representa."),
    ("clave_correcta", "La letra de la opción correcta."),
    ("explicacion", "Explicación general del ítem."),
)


# ------------------------------------------------------------------------------- lectura


def _guard(content: bytes) -> None:
    if len(content) > MAX_FILE_BYTES:
        raise FileTooLargeError("El archivo pesa más de 10 MB.")
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise UnsupportedFileFormatError("El archivo no es una plantilla XLSX.")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        names = {info.filename for info in infos}
        if "[Content_Types].xml" not in names:
            raise UnsupportedFileFormatError("El archivo no es una plantilla XLSX.")
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            raise UnsupportedFileFormatError("No se admiten archivos con macros (.xlsm).")
        if sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_BYTES:
            raise FileTooLargeError("El contenido del archivo es demasiado grande.")
        types = archive.read("[Content_Types].xml")
        if b"macroEnabled" in types:
            raise UnsupportedFileFormatError("No se admiten archivos con macros (.xlsm).")


def _rows(content: bytes) -> Iterator[tuple[int, dict[str, Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as error:  # XML hostil, ZIP dañado, partes faltantes…
        raise UnreadableFileError("No se pudo leer la hoja de cálculo.") from error
    try:
        if SHEET not in workbook.sheetnames:
            raise UnreadableFileError(f"Falta la hoja «{SHEET}» de la plantilla.")
        rows = workbook[SHEET].iter_rows(values_only=True)
        header = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        missing = [name for name in COLUMNS if name not in header]
        if missing:
            raise UnreadableFileError(f"Faltan columnas de la plantilla: {', '.join(missing)}.")
        index = {
            name: header.index(name) for name in (*COLUMNS, *OPTIONAL_COLUMNS) if name in header
        }
        for number, values in enumerate(rows, start=2):
            record = dict.fromkeys(OPTIONAL_COLUMNS)
            record |= {
                name: _clean(values[position] if position < len(values) else None)
                for name, position in index.items()
            }
            if any(value is not None for value in record.values()):
                yield number, record
    finally:
        workbook.close()


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _text(value: Any) -> str | None:
    return None if value is None else str(value)


def read_xlsx(content: bytes) -> Bundle:
    _guard(content)
    bundle = Bundle(schema_version=SCHEMA_VERSION)
    contexts: dict[str, ContextSpec] = {}
    count = 0
    for number, record in _rows(content):
        count += 1
        if count > MAX_IMPORT_ITEMS:
            raise BatchTooLargeError(f"El archivo tiene más de {MAX_IMPORT_ITEMS} ítems.")
        code = _text(record["id"]) or f"fila {number}"
        problem = _unsupported(record)
        if problem is not None:
            bundle.errors.append(EntryError(code, EntryKind.ITEM, (problem,)))
            continue
        context_problem = _context(record, contexts, bundle)
        if context_problem is not None:
            bundle.errors.append(EntryError(code, EntryKind.ITEM, (context_problem,)))
            continue
        item = _item(record, code)
        if isinstance(item, EntryError):
            bundle.errors.append(item)
        else:
            bundle.items.append(item)
    bundle.contexts = list(contexts.values())
    return bundle


def _unsupported(record: dict[str, Any]) -> ContentIssue | None:
    module = _text(record["modulo"])
    if module == "CE":
        return ContentIssue(
            "modulo",
            ValidationRule.UNSUPPORTED_IN_TEMPLATE,
            "Las consignas de Comunicación escrita se importan con el formato JSON.",
        )
    if module == "ING" and record["parte"] == 1:
        return ContentIssue(
            "parte",
            ValidationRule.UNSUPPORTED_IN_TEMPLATE,
            "La parte 1 de Inglés (emparejamiento) se importa con el formato JSON.",
        )
    return None


def _context(
    record: dict[str, Any],
    contexts: dict[str, ContextSpec],
    bundle: Bundle,
) -> ContentIssue | None:
    """Registra el contexto de la fila si es nuevo; `None` si la fila es coherente."""
    context_code = _text(record["contexto_id"])
    text = _text(record["contexto_texto"])
    if context_code is None or (text is None and record["contexto_titulo"] is None):
        return None  # sin contexto, o uno que ya está en el banco
    known = contexts.get(context_code)
    if known is not None:
        if text is not None and text != known.body_md:
            return ContentIssue(
                "contexto_texto",
                ValidationRule.CONTEXT_TEXT_MISMATCH,
                f"El texto del contexto {context_code} no coincide con el de su primera fila.",
            )
        return None
    raw = {
        "id": context_code,
        "modulo": record["modulo"],
        "titulo": _text(record["contexto_titulo"]) or context_code,
        "tipo": "texto",
        "contenido_md": text,
        "tipo_texto": _text(record["contexto_tipo_texto"]),
        "subtipo": _text(record["contexto_subtipo"]),
        "parte": record["parte"],
        "origen": record["origen"] or "docente",
    }
    try:
        contexts[context_code] = mapping.context_spec(
            m.Contexto.model_validate({k: v for k, v in raw.items() if v is not None})
        )
    except ValidationError as error:
        bundle.errors.append(EntryError(context_code, EntryKind.CONTEXT, issues_from(error)))
    return None


def _item(record: dict[str, Any], code: str) -> ItemSpec | EntryError:
    key = _text(record["clave_correcta"])
    options = [
        {
            "clave": letter,
            "texto": _text(record[f"opcion_{letter}"]),
            "correcta": letter == key,
            "retroalimentacion": _text(record[f"retroalimentacion_{letter}"]) or "",
        }
        for letter in KEYS
        if record[f"opcion_{letter}"] is not None
    ]
    raw: dict[str, Any] = {
        "id": _text(record["id"]),
        "modulo": record["modulo"],
        "parte": record["parte"],
        **{tag: _text(record[tag]) for tag in _TAGS},
        "dificultad": record["dificultad"],
        "origen": record["origen"] or "docente",
        "contexto_id": _text(record["contexto_id"]),
        "enunciado": _text(record["enunciado"]),
        "opciones": options,
        "explicacion": _text(record["explicacion"]) or "",
    }
    try:
        return mapping.item_spec(
            m.Item.model_validate({k: v for k, v in raw.items() if v is not None})
        )
    except ValidationError as error:
        return EntryError(code, EntryKind.ITEM, issues_from(error))


# ------------------------------------------------------------------------------- plantilla

_LISTS: tuple[tuple[str, str, TaxonomyNodeKind | None], ...] = (
    ("competencias", "competencia", TaxonomyNodeKind.COMPETENCE),
    ("afirmaciones", "afirmacion", TaxonomyNodeKind.ASSERTION),
    ("evidencias", "evidencia", TaxonomyNodeKind.EVIDENCE),
    ("componentes", "componente", TaxonomyNodeKind.COMPONENT),
    ("situaciones", "situacion", TaxonomyNodeKind.SITUATION),
    ("tipos_texto", "contexto_tipo_texto", TaxonomyNodeKind.TEXT_TYPE),
)
_FIXED: tuple[tuple[str, str, tuple[str, ...]], ...] = (
    ("modulos", "modulo", ("LC", "RC", "CC", "ING")),
    ("dificultades", "dificultad", ("baja", "media", "alta")),
    ("origenes", "origen", ("docente", "ia")),
    ("claves", "clave_correcta", tuple(KEYS)),
)


def write_template(nodes: Iterable[TaxonomyNode]) -> bytes:
    """Plantilla con las categorías activas en la hoja `listas` y validaciones de celda."""
    active = [node for node in nodes if node.active]
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    items = workbook.create_sheet(SHEET)
    _items_header(items)
    items.freeze_panes = "B2"
    _instructions(workbook)

    lists = workbook.create_sheet("listas")
    columns: list[tuple[str, str, list[str]]] = [
        (title, target, list(values)) for title, target, values in _FIXED
    ]
    for title, target, kind in _LISTS:
        columns.append((title, target, [n.code for n in active if n.kind is kind]))
    for position, (title, target, values) in enumerate(columns, start=1):
        letter = get_column_letter(position)
        lists.cell(row=1, column=position, value=title)
        for line, value in enumerate(values, start=2):
            lists.cell(row=line, column=position, value=value)
        if values:
            _validate(items, target, f"listas!${letter}$2:${letter}${len(values) + 1}")

    parts = DataValidation(type="whole", operator="between", formula1="2", formula2="7")
    parts.error = "La parte debe ser un número de 2 a 7."
    items.add_data_validation(parts)
    parts.add(_range("parte"))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _items_header(sheet: Any) -> None:
    for position, name in enumerate(_ALL_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = (
            40 if name in ("enunciado", "contexto_texto", "explicacion") else 18
        )
    sheet.append(list(_ALL_COLUMNS))


def _instructions(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("instrucciones")
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 100
    sheet.append(["Columna", "Qué escribir"])
    for row in _INSTRUCTIONS:
        sheet.append(list(row))


def _range(column: str) -> str:
    letter = get_column_letter(COLUMNS.index(column) + 1)
    return f"{letter}2:{letter}{TEMPLATE_ROWS + 1}"


def _validate(sheet: Any, column: str, source: str) -> None:
    validation = DataValidation(type="list", formula1=source, allow_blank=True)
    validation.error = "Elige un valor de la lista."
    sheet.add_data_validation(validation)
    validation.add(_range(column))


# ------------------------------------------------------------------------------- exportación

_ALL_COLUMNS = (*COLUMNS, *OPTIONAL_COLUMNS)
_TEMPLATE_MODULES = frozenset({"LC", "RC", "CC", "ING"})
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class XlsxExportWriter:
    """Exportación en la plantilla (FR-038): solo lo que la plantilla puede representar.

    Quedan fuera las consignas, los ítems de emparejamiento (Inglés parte 1) y los que usan un
    contexto que no es de texto (imagen, tabla u opciones compartidas).
    """

    media_type = XLSX_MEDIA_TYPE
    extension = "xlsx"
    streams = False

    def __init__(self) -> None:
        self._workbook = Workbook(write_only=True)
        self._sheet = self._workbook.create_sheet(SHEET)
        _items_header(self._sheet)

    def accepts(self, item: ItemSpec, context: ContextSpec | None) -> bool:
        return (
            item.kind is ItemKind.SELECTION
            and item.module_code in _TEMPLATE_MODULES
            and all(option.key in KEYS for option in item.options or ())
            and (
                context is None
                or (context.kind is ContextKind.TEXT and context.shared_options is None)
            )
        )

    def begin(self) -> bytes:
        return b""

    def items(self, entries: Iterable[tuple[ItemSpec, ContextSpec | None]]) -> bytes:
        for item, context in entries:
            self._sheet.append([self._cell(value) for value in _row(item, context)])
        return b""

    def end(self, tail: Bundle) -> bytes:  # noqa: ARG002 - cada fila ya lleva su contexto
        _instructions(self._workbook)
        buffer = io.BytesIO()
        self._workbook.save(buffer)
        return buffer.getvalue()

    def _cell(self, value: Any) -> Any:
        cell = WriteOnlyCell(self._sheet, value=value)
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"  # texto, nunca una fórmula que Excel evaluaría
        return cell


def _row(item: ItemSpec, context: ContextSpec | None) -> list[Any]:
    values: dict[str, Any] = mapping.item_out(item)
    if context is not None:
        values["contexto_titulo"] = context.title
        values["contexto_tipo_texto"] = context.text_type_code
        values["contexto_texto"] = context.body_md
        values["contexto_subtipo"] = context.subtype
    for option in item.options or ():
        values[f"opcion_{option.key}"] = option.text
        values[f"retroalimentacion_{option.key}"] = option.feedback
        if option.correct:
            values["clave_correcta"] = option.key
    return [values.get(name) for name in _ALL_COLUMNS]


# ------------------------------------------------------------------------------- reporte

_REPORT_HEADER = ("Código", "Entidad", "Campo", "Regla", "Mensaje")


def write_report(rows: Iterable[tuple[str, str, str, str, str]]) -> bytes:
    """Reporte de errores de una importación (FR-033): una fila por problema."""
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    sheet = workbook.create_sheet("errores")
    sheet.append(list(_REPORT_HEADER))
    for row in rows:
        sheet.append(list(row))
    for letter, width in zip("ABCDE", (16, 14, 28, 26, 80), strict=True):
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
