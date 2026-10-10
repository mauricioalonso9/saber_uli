"""T085: exportación del banco (US6; FR-038, SC-008; research R-03, R-04 y R-18).

Se exporta en el formato de importación (JSON v1 o la plantilla XLSX), con filtros y la versión
publicada o la última de cada ítem. Reimportar lo exportado no cambia nada. La plantilla solo
lleva lo que puede representar y cuenta el resto. Cada exportación queda registrada.
"""

import io
import json
from datetime import UTC, datetime
from typing import Any
from uuid import uuid4

import httpx
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine

from saber_uli.content.application.bank_export import (
    EXPORT_BATCH_SIZE,
    BankExport,
    ExportFile,
    ExportRequest,
)
from saber_uli.content.domain.item_version import VersionState
from saber_uli.content.infrastructure.interchange.export_format import EXPORT_WRITERS
from saber_uli.content.infrastructure.interchange.xlsx_format import COLUMNS
from saber_uli.content.infrastructure.unit_of_work import SqlAlchemyContentUnitOfWork
from saber_uli.identity.domain.roles import Role
from saber_uli.shared.application.event_bus import EventBus
from saber_uli.shared.domain.clock import FixedClock
from saber_uli.shared.infrastructure.db import create_session_factory
from tests.integration.content.conftest import UowFactory
from tests.integration.content.imports import importer
from tests.integration.content.seed_bank import bank, publish, seed, seed_codes
from tests.integration.identity.staff import StaffFactory

__all__ = ["bank"]  # fixture

NOW = datetime(2026, 10, 9, 21, 0, tzinfo=UTC)
TEACHER = uuid4()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ING_PART_1 = [f"ING-{n:04d}" for n in range(1, 6)]
TABLE_ITEMS = ["LC-0007", "LC-0008"]  # su contexto LC-CTX-003 es una tabla


def exporter(uow: UowFactory, *, batch_size: int = EXPORT_BATCH_SIZE) -> BankExport:
    return BankExport(
        uow_factory=uow, clock=FixedClock(NOW), writers=EXPORT_WRITERS, batch_size=batch_size
    )


async def content_of(file: ExportFile) -> bytes:
    return b"".join([chunk async for chunk in file.chunks])


async def export_json(uow: UowFactory, **filters: Any) -> dict[str, Any]:
    file = await exporter(uow).export(ExportRequest(format="json", **filters), TEACHER)
    data: dict[str, Any] = json.loads(await content_of(file))
    return data


def codes(data: dict[str, Any]) -> list[str]:
    return sorted(e["id"] for e in data["items"] + data["consignas_escritura"])


async def node_id(uow: UowFactory, code: str) -> Any:
    async with uow() as unit:
        return next(node.id for node in await unit.taxonomy.list() if node.code == code)


async def test_json_con_los_filtros(content_uow: UowFactory, bank: None) -> None:
    published = ["CE-0001", "ING-0006", "ING-0008", "LC-0001", "LC-0002", "RC-0001", "RC-0003"]
    await publish(content_uow, *published)

    assert codes(await export_json(content_uow)) == published
    assert codes(await export_json(content_uow, module="RC")) == ["RC-0001", "RC-0003"]
    low = await export_json(content_uow, difficulty="low")
    assert codes(low) == ["ING-0006", "LC-0001", "RC-0001"]
    assert {e["dificultad"] for e in low["items"]} == {"baja"}
    rc_c1 = await node_id(content_uow, "RC-C1")
    assert codes(await export_json(content_uow, taxonomy_node=rc_c1)) == ["RC-0001"]
    in_review = await export_json(content_uow, version="latest", state="in_review", module="LC")
    assert codes(in_review) == [f"LC-{n:04d}" for n in range(3, 9)]
    assert codes(await export_json(content_uow, origin="teacher")) == []
    assert codes(await export_json(content_uow, version="latest", origin="ai")) == seed_codes()


async def test_la_publicada_frente_a_la_ultima(content_uow: UowFactory, bank: None) -> None:
    await publish(content_uow, "LC-0001")
    [entry] = (await export_json(content_uow, module="LC"))["items"]
    changed = {**entry, "enunciado": "¿Qué propone el autor del texto?"}
    service = importer(content_uow, NOW)
    job = await service.preview(
        json.dumps({"esquema_version": "1.0", "items": [changed]}).encode(), "b.json", TEACHER
    )
    assert job.id is not None
    await service.confirm(
        job.id, TEACHER, target_state=VersionState.DRAFT, originality_declared=True
    )

    [published] = (await export_json(content_uow, module="LC"))["items"]
    latest = await export_json(content_uow, module="LC", version="latest")
    [newest] = [e for e in latest["items"] if e["id"] == "LC-0001"]

    assert (published["version"], published["estado"]) == (1, "publicado")
    assert published["enunciado"] == entry["enunciado"]
    assert (newest["version"], newest["estado"]) == (2, "en_revision")  # origen IA
    assert newest["enunciado"] == "¿Qué propone el autor del texto?"


async def test_incluye_contextos_categorias_y_rubricas(content_uow: UowFactory, bank: None) -> None:
    await publish(content_uow, "LC-0004", "CE-0002")

    data = await export_json(content_uow)

    assert [c["id"] for c in data["contextos"]] == ["LC-CTX-002"]
    original = next(c for c in seed()["contextos"] if c["id"] == "LC-CTX-002")
    assert data["contextos"][0]["contenido_md"] == original["contenido_md"]
    assert [r["id"] for r in data["rubricas"]] == ["RUB-CE-01"]
    modules = {m["codigo"]: m for m in data["modulos"]}
    assert set(modules) == {"CE", "LC"}
    item = data["items"][0]
    [competence] = modules["LC"]["competencias"]
    assert competence["codigo"] == item["competencia"]
    [assertion] = competence["afirmaciones"]
    assert assertion["codigo"] == item["afirmacion"]
    assert [e["codigo"] for e in assertion["evidencias"]] == [item["evidencia"]]
    assert list(modules["LC"]["tipos_texto"]) == [original["tipo_texto"]]


async def test_reimportar_lo_exportado_no_cambia_nada(content_uow: UowFactory, bank: None) -> None:
    await publish(content_uow, *seed_codes())
    file = await exporter(content_uow).export(ExportRequest(format="json"), TEACHER)

    job = await importer(content_uow, NOW).preview(await content_of(file), file.filename, TEACHER)

    assert job.counts() == {"create": 0, "update": 0, "unchanged": 54, "error": 0}


async def test_xlsx_solo_con_lo_que_la_plantilla_representa(
    content_uow: UowFactory, bank: None
) -> None:
    await publish(content_uow, *seed_codes())

    file = await exporter(content_uow).export(ExportRequest(format="xlsx"), TEACHER)
    content = await content_of(file)

    assert file.media_type == XLSX
    assert file.filename.endswith(".xlsx")
    sheet = load_workbook(io.BytesIO(content), read_only=True)["items"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][: len(COLUMNS)] == COLUMNS
    outside = {*ING_PART_1, *TABLE_ITEMS}
    expected = [c for c in seed_codes() if c not in outside and not c.startswith("CE-")]
    assert sorted(str(row[0]) for row in rows[1:]) == expected
    assert file.skipped == len(outside) + 3  # emparejamiento, tabla y consignas
    job = await importer(content_uow, NOW).preview(content, file.filename, TEACHER)
    assert job.counts()["create"] + job.counts()["update"] + job.counts()["error"] == 0


async def test_cada_exportacion_queda_registrada(
    content_uow: UowFactory, bank: None, app_engine: AsyncEngine
) -> None:
    await publish(content_uow, *seed_codes("ING"))

    file = await exporter(content_uow).export(
        ExportRequest(format="xlsx", module="ING", version="published"), TEACHER
    )
    await content_of(file)

    async with app_engine.connect() as conn:
        row = (
            await conn.execute(
                text(
                    "SELECT created_by, created_at, format, filter, version_choice, item_count,"
                    " skipped_count FROM content.exports"
                )
            )
        ).one()
    assert row.created_by == TEACHER
    assert row.created_at == NOW
    assert (row.format, row.version_choice) == ("xlsx", "published")
    assert row.filter == {"module": "ING"}
    assert (row.item_count, row.skipped_count) == (14, 5)


async def test_lee_por_bloques(
    content_uow: UowFactory, bank: None, app_engine: AsyncEngine
) -> None:
    assert EXPORT_BATCH_SIZE == 500
    limits: list[int] = []
    factory = create_session_factory(app_engine)

    class CountingItems:
        def __init__(self, inner: Any) -> None:
            self._inner = inner

        async def export_batch(self, *args: Any, **kwargs: Any) -> Any:
            limits.append(kwargs["limit"])
            return await self._inner.export_batch(*args, **kwargs)

        def __getattr__(self, name: str) -> Any:
            return getattr(self._inner, name)

    class CountingUow(SqlAlchemyContentUnitOfWork):
        @property
        def items(self) -> Any:
            return CountingItems(super().items)

    service = exporter(lambda: CountingUow(factory, EventBus()), batch_size=10)
    file = await service.export(ExportRequest(format="json", version="latest"), TEACHER)
    data = json.loads(await content_of(file))

    assert codes(data) == seed_codes()
    assert limits == [10] * 5 + [10] * (len(limits) - 5)
    assert len(limits) in (5, 6)  # 46 entradas en bloques de 10, más el bloque vacío final


async def test_api_exporta_con_encabezados(
    api_client: httpx.AsyncClient, staff: StaffFactory, content_uow: UowFactory, bank: None
) -> None:
    teacher = await staff(Role.TEACHER)
    await publish(content_uow, *seed_codes("ING"))

    xlsx = await api_client.get(
        "/api/v1/content/export",
        params={"format": "xlsx", "module": "ING"},
        headers=teacher.headers,
    )
    json_file = await api_client.get(
        "/api/v1/content/export", params={"format": "json"}, headers=teacher.headers
    )
    student = await api_client.get(
        "/api/v1/content/export", params={"format": "json"}, headers=(await staff()).headers
    )

    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.headers["content-type"] == XLSX
    assert xlsx.headers["x-skipped-items"] == "5"
    assert xlsx.headers["content-disposition"].startswith("attachment;")
    assert ".xlsx" in xlsx.headers["content-disposition"]
    assert json_file.status_code == 200
    assert json_file.headers["content-type"].startswith("application/json")
    assert json_file.headers["x-skipped-items"] == "0"
    assert len(json.loads(json_file.content)["items"]) == 19  # el JSON representa todo
    assert student.status_code == 403
