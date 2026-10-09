"""T068: API de importación (US4; contrato: `createImport`, `getImport`, `getImportReport`,
`confirmImport`, `listImports`, `getImportTemplate`; FR-028 a FR-037; research R-05).
"""

import io
import logging

import httpx
import pytest
from openpyxl import load_workbook

from saber_uli.content.application.bundle import MAX_IMPORT_BYTES
from saber_uli.identity.domain.roles import Role
from tests.integration.content.conftest import LcTaxonomy
from tests.integration.content.imports import bank_file, lc_item
from tests.integration.identity.staff import Staff, StaffFactory

URL = "/api/v1/content/imports"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
async def people(staff: StaffFactory) -> dict[str, Staff]:
    return {
        "teacher": await staff(Role.TEACHER),
        "other": await staff(Role.TEACHER),
        "admin": await staff(Role.ADMIN),
    }


def upload(content: bytes, name: str = "banco-secreto.json") -> dict[str, tuple[str, bytes, str]]:
    return {"file": (name, content, "application/octet-stream")}


async def test_vista_previa_filtros_reporte_y_confirmacion(
    api_client: httpx.AsyncClient,
    people: dict[str, Staff],
    lc_taxonomy: LcTaxonomy,
    caplog: pytest.LogCaptureFixture,
) -> None:
    teacher = people["teacher"].headers
    content = bank_file([lc_item(n) for n in range(1, 4)] + [lc_item(4, afirmacion="LC-A9")])
    caplog.set_level(logging.INFO)

    created = await api_client.post(URL, files=upload(content), headers=teacher)

    assert created.status_code == 201, created.text
    body = created.json()
    assert body["status"] == "pending"
    assert body["source_name"] == "banco-secreto.json"
    assert body["format"] == "json"
    assert body["counts"] == {"create": 4, "update": 0, "unchanged": 0, "error": 1}
    assert body["created_by"]["id"] == str(people["teacher"].id)
    assert body["result"] is None
    assert "banco-secreto" not in caplog.text

    detail = f"{URL}/{body['id']}"
    errors = await api_client.get(detail, params={"action": "error"}, headers=teacher)
    page = await api_client.get(detail, params={"page": 2, "page_size": 2}, headers=teacher)
    assert [row["code"] for row in errors.json()["rows"]] == ["LC-0004"]
    assert errors.json()["rows"][0]["errors"][0]["rule"] == "taxonomy_unknown"
    assert page.json()["meta"] == {"page": 2, "page_size": 2, "total": 5}
    assert len(page.json()["rows"]) == 2

    report = await api_client.get(f"{detail}/report", headers=teacher)
    assert report.status_code == 200
    assert report.headers["content-type"] == XLSX
    assert "attachment" in report.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(report.content)).active
    assert sheet is not None
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[0] == ["Código", "Entidad", "Campo", "Regla", "Mensaje"]
    assert values[1][:4] == ["LC-0004", "item", "taxonomy.assertion_id", "taxonomy_unknown"]

    missing = await api_client.post(
        f"{detail}/confirmation", json={"target_state": "draft"}, headers=teacher
    )
    by_other = await api_client.post(
        f"{detail}/confirmation",
        json={"target_state": "draft", "originality_declared": True},
        headers=people["other"].headers,
    )
    done = await api_client.post(
        f"{detail}/confirmation",
        json={"target_state": "in_review", "originality_declared": True},
        headers=teacher,
    )
    again = await api_client.post(
        f"{detail}/confirmation",
        json={"target_state": "in_review", "originality_declared": True},
        headers=teacher,
    )

    assert missing.status_code == 422
    assert by_other.status_code == 404
    assert done.status_code == 200, done.text
    assert done.json()["status"] == "confirmed"
    assert done.json()["result"] == {"created": 3, "versioned": 0, "skipped": []}
    assert again.status_code == 409
    assert again.json()["type"].endswith("import-not-pending")


async def test_rechazos_del_archivo(
    api_client: httpx.AsyncClient, people: dict[str, Staff], lc_taxonomy: LcTaxonomy
) -> None:
    teacher = people["teacher"].headers

    big = await api_client.post(URL, files=upload(b" " * (MAX_IMPORT_BYTES + 1)), headers=teacher)
    many = await api_client.post(
        URL, files=upload(bank_file([lc_item(n) for n in range(1, 1002)])), headers=teacher
    )
    other = await api_client.post(URL, files=upload(b"a;b", "banco.csv"), headers=teacher)
    unreadable = await api_client.post(URL, files=upload(b"{roto"), headers=teacher)

    assert (big.status_code, big.json()["type"].split(":")[-1]) == (413, "file-too-large")
    assert (many.status_code, many.json()["type"].split(":")[-1]) == (413, "batch-too-large")
    assert (other.status_code, other.json()["type"].split(":")[-1]) == (
        415,
        "unsupported-file-format",
    )
    assert (unreadable.status_code, unreadable.json()["type"].split(":")[-1]) == (
        422,
        "unreadable-file",
    )


async def test_cada_docente_ve_sus_importaciones_y_el_administrador_todas(
    api_client: httpx.AsyncClient, people: dict[str, Staff], lc_taxonomy: LcTaxonomy
) -> None:
    mine = await api_client.post(
        URL, files=upload(bank_file([lc_item(1)])), headers=people["teacher"].headers
    )
    theirs = await api_client.post(
        URL, files=upload(bank_file([lc_item(2)])), headers=people["other"].headers
    )

    own = await api_client.get(URL, headers=people["teacher"].headers)
    everything = await api_client.get(URL, headers=people["admin"].headers)
    foreign = await api_client.get(
        f"{URL}/{theirs.json()['id']}", headers=people["teacher"].headers
    )

    assert [i["id"] for i in own.json()["items"]] == [mine.json()["id"]]
    assert {i["id"] for i in everything.json()["items"]} == {
        mine.json()["id"],
        theirs.json()["id"],
    }
    assert everything.json()["meta"]["total"] == 2
    assert foreign.status_code == 404


async def test_plantilla(
    api_client: httpx.AsyncClient, people: dict[str, Staff], lc_taxonomy: LcTaxonomy
) -> None:
    response = await api_client.get(
        "/api/v1/content/import-template", headers=people["teacher"].headers
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert 'filename="plantilla-banco-saber-uli.xlsx"' in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["items", "instrucciones", "listas"]


async def test_limite_de_30_subidas_por_hora(
    api_client: httpx.AsyncClient, people: dict[str, Staff], lc_taxonomy: LcTaxonomy
) -> None:
    teacher = people["teacher"].headers
    statuses = [
        (await api_client.post(URL, files=upload(b"{roto"), headers=teacher)).status_code
        for _ in range(31)
    ]
    blocked = await api_client.post(URL, files=upload(b"{roto"), headers=teacher)

    assert statuses == [422] * 30 + [429]
    assert int(blocked.headers["Retry-After"]) > 0
