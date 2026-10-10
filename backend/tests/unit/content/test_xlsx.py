"""T065: plantilla XLSX de ítems de selección (US4; FR-029, FR-037; research R-04).

La plantilla tiene las hojas `items`, `instrucciones` y `listas`, con validaciones de celda que
usan los códigos activos. Una plantilla llena se lee como el formato v1. Las defensas actúan
antes de abrir el archivo: tamaño, bomba zip, `.xlsm` y ZIP sin `[Content_Types].xml`. Las
fórmulas se leen como texto, sin evaluarse, y el XML con entidades externas se rechaza.
"""

import io
import zipfile
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from saber_uli.content.application.bundle import EntryKind, UnreadableFileError
from saber_uli.content.domain.context import ContextKind
from saber_uli.content.domain.errors import FileTooLargeError
from saber_uli.content.domain.item_version import Difficulty, ItemKind
from saber_uli.content.domain.taxonomy import TaxonomyNode, TaxonomyNodeKind
from saber_uli.content.domain.validation import ValidationRule
from saber_uli.content.infrastructure.interchange.xlsx_format import (
    COLUMNS,
    MAX_FILE_BYTES,
    MAX_UNCOMPRESSED_BYTES,
    OPTIONAL_COLUMNS,
    UnsupportedFileFormatError,
    read_xlsx,
    write_report,
    write_template,
)

K = TaxonomyNodeKind


def _node(module: str, kind: TaxonomyNodeKind, code: str, *, active: bool = True) -> TaxonomyNode:
    return TaxonomyNode(uuid4(), module, kind, code, code, 1, active=active)


NODES = [
    _node("LC", K.COMPETENCE, "LC-C1"),
    _node("LC", K.ASSERTION, "LC-A1"),
    _node("LC", K.EVIDENCE, "LC-E1"),
    _node("LC", K.EVIDENCE, "LC-E9", active=False),
    _node("LC", K.TEXT_TYPE, "LC-TT-CI"),
    _node("ING", K.PART, "ING-P2"),
]


def row(**values: object) -> dict[str, object]:
    base: dict[str, object] = {
        "id": "LC-0901",
        "modulo": "LC",
        "competencia": "LC-C1",
        "afirmacion": "LC-A1",
        "evidencia": "LC-E1",
        "dificultad": "media",
        "origen": "docente",
        "contexto_id": "LC-CTX-901",
        "contexto_titulo": "Desconectar para aprender",
        "contexto_tipo_texto": "LC-TT-CI",
        "contexto_texto": "Un texto breve de opinión.",
        "enunciado": "¿Cuál es la tesis del autor?",
        "clave_correcta": "A",
        "explicacion": "Está en el primer párrafo.",
    }
    for key in "ABCD":
        base[f"opcion_{key}"] = f"Opción {key}"
        base[f"retroalimentacion_{key}"] = f"Porque {key}."
    base.update(values)
    return base


def filled(*rows: dict[str, object]) -> bytes:
    workbook = load_workbook(io.BytesIO(write_template(NODES)))
    sheet = workbook["items"]
    for index, values in enumerate(rows, start=2):
        for column, name in enumerate(COLUMNS, start=1):
            sheet.cell(row=index, column=column, value=values.get(name))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_la_plantilla_tiene_hojas_y_validaciones_con_codigos_activos() -> None:
    workbook = load_workbook(io.BytesIO(write_template(NODES)))

    assert workbook.sheetnames == ["items", "instrucciones", "listas"]
    assert [cell.value for cell in workbook["items"][1]] == [*COLUMNS, *OPTIONAL_COLUMNS]
    listed = {str(cell.value) for column in workbook["listas"].iter_cols() for cell in column[1:]}
    assert {"LC-C1", "LC-A1", "LC-E1", "LC-TT-CI"} <= listed
    assert "LC-E9" not in listed
    formulas = {v.formula1 for v in workbook["items"].data_validations.dataValidation}
    assert any("listas!" in formula for formula in formulas)


def test_una_plantilla_llena_se_lee_como_el_formato_v1() -> None:
    bundle = read_xlsx(filled(row(), row(id="LC-0902", enunciado="¿Qué propone el autor?")))

    assert bundle.errors == []
    [context] = bundle.contexts
    assert (context.code, context.kind, context.text_type_code) == (
        "LC-CTX-901",
        ContextKind.TEXT,
        "LC-TT-CI",
    )
    first, second = bundle.items
    assert (first.code, first.kind, first.difficulty) == (
        "LC-0901",
        ItemKind.SELECTION,
        Difficulty.MEDIUM,
    )
    assert first.tags[K.ASSERTION] == "LC-A1"
    assert first.context_code == second.context_code == "LC-CTX-901"
    assert [o.correct for o in first.options or ()] == [True, False, False, False]
    assert first.options and first.options[1].feedback == "Porque B."


def test_un_texto_distinto_para_el_mismo_contexto_es_un_error() -> None:
    bundle = read_xlsx(filled(row(), row(id="LC-0902", contexto_texto="Otro texto.")))

    [error] = bundle.errors
    assert error.code == "LC-0902"
    assert [(i.field, i.rule) for i in error.issues] == [
        ("contexto_texto", ValidationRule.CONTEXT_TEXT_MISMATCH)
    ]
    assert [item.code for item in bundle.items] == ["LC-0901"]


@pytest.mark.parametrize(
    "values",
    [
        {"id": "ING-0901", "modulo": "ING", "parte": 1, "competencia": None},
        {"id": "CE-0901", "modulo": "CE", "competencia": None},
    ],
)
def test_ingles_parte_1_y_ce_no_van_en_la_plantilla(values: dict[str, object]) -> None:
    bundle = read_xlsx(filled(row(**values)))

    [error] = bundle.errors
    assert error.entity is EntryKind.ITEM
    assert error.issues[0].rule is ValidationRule.UNSUPPORTED_IN_TEMPLATE
    assert bundle.items == []


def test_una_fila_sin_codigo_se_identifica_por_su_numero() -> None:
    bundle = read_xlsx(filled(row(), row(id=None, dificultad="dificilisima")))

    assert bundle.errors[0].code == "fila 3"


def test_las_formulas_se_leen_como_texto() -> None:
    bundle = read_xlsx(filled(row(enunciado='=HYPERLINK("http://x","clic")')))

    assert bundle.items[0].stem_md == '=HYPERLINK("http://x","clic")'


def test_sin_la_hoja_items_o_sin_columnas_el_archivo_es_ilegible() -> None:
    workbook = Workbook()
    workbook.active.title = "otra"  # type: ignore[union-attr]
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(UnreadableFileError):
        read_xlsx(buffer.getvalue())


def _zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def test_defensas_antes_de_abrir() -> None:
    with pytest.raises(FileTooLargeError):
        read_xlsx(b"PK" + b"\0" * MAX_FILE_BYTES)
    with pytest.raises(UnsupportedFileFormatError):
        read_xlsx(b"no es un zip")
    with pytest.raises(UnsupportedFileFormatError):
        read_xlsx(_zip({"xl/workbook.xml": b"<x/>"}))
    with pytest.raises(UnsupportedFileFormatError):
        macro = filled(row())
        with zipfile.ZipFile(io.BytesIO(macro)) as source:
            files = {name: source.read(name) for name in source.namelist()}
        read_xlsx(_zip(files | {"xl/vbaProject.bin": b"macro"}))
    bomb = _zip(
        {"[Content_Types].xml": b"<Types/>", "xl/big.xml": b"0" * (MAX_UNCOMPRESSED_BYTES + 1)}
    )
    assert len(bomb) < MAX_FILE_BYTES
    with pytest.raises(FileTooLargeError):
        read_xlsx(bomb)


def test_el_xml_con_entidades_externas_se_rechaza() -> None:
    template = filled(row())
    with zipfile.ZipFile(io.BytesIO(template)) as source:
        files = {name: source.read(name) for name in source.namelist()}
    files["xl/workbook.xml"] = (
        b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY secreto SYSTEM "file:///etc/passwd">]>'
        b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"&secreto;</workbook>"
    )

    with pytest.raises(UnreadableFileError):
        read_xlsx(_zip(files))


def test_el_reporte_escribe_como_texto_lo_que_empieza_por_igual() -> None:
    """T115 (ASVS 5.3.1): un código o mensaje del archivo nunca llega como fórmula."""
    hostile = '=HYPERLINK("http://ev.il","x")'
    content = write_report([(hostile, "item", "code", "invalid_value", "=1+1")])

    sheet = load_workbook(io.BytesIO(content))["errores"]

    assert sheet["A2"].value == hostile
    assert sheet["A2"].data_type == "s"
    assert sheet["E2"].data_type == "s"
